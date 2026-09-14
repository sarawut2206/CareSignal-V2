# -*- coding: utf-8 -*-
"""
make-bangkhae-dayplan.py — แผนวันลงพื้นที่บ้านบางแค 1 หน้า A4 ไว้พกไปหน้างาน
python tools/make-bangkhae-dayplan.py   → validation/แผนลงพื้นที่-บ้านบางแค.docx และ .pdf

ใช้ตัวช่วยจัดรูปแบบชุดเดียวกับ make-bangkhae-kit.py ให้ฟอนต์ สี และการตัดคำไทยตรงกัน
"""
import os, sys, importlib.util
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location("kit", os.path.join(HERE, "make-bangkhae-kit.py"))
kit = importlib.util.module_from_spec(spec); spec.loader.exec_module(kit)
from_kit = ("para", "rich", "table", "box", "borders", "shade", "cell_text", "set_font",
            "BRAND", "INK2", "WARN", "STOP", "FONT", "SCRIPT_30S", "ROOT", "OUT_DIR")
globals().update({k: getattr(kit, k) for k in from_kit})

OUT_DOCX = os.path.join(OUT_DIR, "แผนลงพื้นที่-บ้านบางแค.docx")
OUT_PDF = os.path.join(OUT_DIR, "แผนลงพื้นที่-บ้านบางแค.pdf")
S = 12.5   # ขนาดเนื้อความ


def head(where, text):
    return para(where, text, size=14, bold=True, color=BRAND, before=3, after=0, keep=True)


def item(where, text, mark="☐ "):
    return rich(where, [(mark, True, BRAND), (text, False, None)], size=S, after=0, indent=0.15)


def two_col(doc, left, right, widths=(8.5, 8.5)):
    t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t, color="C9DAF3")
    for i, fill in enumerate((left, right)):
        c = t.rows[0].cells[i]; c.width = Cm(widths[i])
        blank = c.paragraphs[0]._element
        fill(c)
        blank.getparent().remove(blank)   # ย่อหน้าว่างที่ Word ใส่มากับเซลล์ กินหนึ่งบรรทัดบนสุด
    para(doc, "", size=3, after=0)


def tonight(c):
    head(c, "คืนนี้ 14 ก.ย.")
    item(c, "เติมหนังสือ (ที่อยู่ ชื่อ 2 จุด เวลา เบอร์) แล้วเซ็น")
    item(c, "พิมพ์: หนังสือ 2 · ประกาศ 2 · ใบยินยอม 6 · ผู้สูงอายุ 4 ชุด · "
            "ครอบครัว 2 · เจ้าหน้าที่ 3 · บัตรฟีเจอร์ 1 · สรุปรายคน 6")
    item(c, "บัตรประชาชน ปากกา คลิปบอร์ด พาวเวอร์แบงก์")
    item(c, "ซ้อมอ่านประโยคแนะนำออกเสียง 3 รอบ")


def morning(c):
    head(c, "เช้า 15 ก.ย. โทรหาศูนย์ฯ ก่อนออกจากบ้าน")
    item(c, "วันนี้เข้าพบได้ไหม ติดต่อใคร ห้องไหน")
    item(c, "ช่วงไหนไม่ควรรบกวน (เวลาอาหาร เวลาพัก)")
    item(c, "ถ่ายภาพได้หรือไม่")
    item(c, "ถ้า ผอ. ต้องอนุมัติก่อน ส่งหนังสือทางไลน์ได้ไหม")
    para(c, "ไปสองคน: คนหนึ่งถาม อีกคนจดคำพูดจริง", size=12, color=INK2, after=1)


def evening(c):
    head(c, "เย็น 15 ก.ย.")
    item(c, "กรอกแบบสรุปรายคนที่ค้างให้ครบ")
    item(c, "ถ่ายรูปแบบสัมภาษณ์ + แบบสรุป ส่งให้ Claude")
    item(c, "ส่งแบบสอบถามหน้า 9 ทาง LINE ให้ครอบครัว 2–3 คน")


def deadline(c):
    head(c, "16 ก.ย. วันส่งภารกิจที่ 3")
    item(c, "เช้า: โทรสัมภาษณ์ครอบครัวถ้ายังไม่ครบ เก็บคำตอบจาก LINE")
    item(c, "บ่าย: ทำใบสรุป 5 หัวข้อ ตรวจว่าตรงกับที่คุยจริง")
    item(c, "ส่งไม่เกิน 20:00 น. อย่ารอ 23:59 น.")


def tighten(doc):
    """TH Sarabun New มีระยะบรรทัดในตัวสูงมาก ระยะ 1 เท่าจึงห่างเกือบ 1.5 เท่าของตัวอักษร
    ตั้งระยะบรรทัดแบบคงที่ตามตัวอักษรที่ใหญ่สุดในย่อหน้า ให้ทั้งแผนลงหนึ่งหน้า"""
    from docx.enum.text import WD_LINE_SPACING
    def walk(paragraphs):
        for p in paragraphs:
            sizes = [r.font.size.pt for r in p.runs if r.font.size]
            size = max(sizes) if sizes else 3
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(max(4, size * 1.3))
    walk(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                walk(c.paragraphs)


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.2), Cm(1.0)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(S)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    rich(doc, [("แผนลงพื้นที่บ้านบางแค", True, BRAND), ("   อังคารที่ 15 กันยายน 2569 · CareSignal ภารกิจที่ 3", False, INK2)], size=19, after=1)
    box(doc, "เป้าหมายวันนี้",
        "เจ้าหน้าที่ 2 คน + ผู้สูงอายุ 3 คน · ทุกคนเซ็นใบยินยอม · แบบสรุปรายคนเสร็จภายในคืนนี้ · "
        "ได้กี่คนเขียนตามจริง ห้ามกรอกแทน")

    two_col(doc, tonight, morning)

    head(doc, "ที่ศูนย์ฯ (รวมประมาณ 3 ชั่วโมง)")
    table(doc, ["", "ทำอะไร", "เวลา"], [
        ["1", "พบผู้ประสานงาน ยื่นหนังสือ + ประกาศ เล่าว่าเป็นใคร ทำอะไร ขอให้แนะนำผู้สูงอายุที่สื่อสารได้ชัด 3 ท่าน", "15 นาที"],
        ["2", "สัมภาษณ์เจ้าหน้าที่ก่อน 2 คน (หน้า 7) เพราะอาจติดงานช่วงหลัง และช่วยชี้ผู้สูงอายุที่เหมาะได้", "25 นาที × 2"],
        ["3", "สัมภาษณ์ผู้สูงอายุทีละคน (หน้า 4–5) ถ้าท่านเหนื่อยให้หยุด", "20 นาที × 3"],
        ["4", "ระหว่างคน กรอกแบบสรุปรายคน (หน้า 10) ทันทีขณะยังจำได้", "5 นาที × 5"],
        ["5", "ถ้ามีญาติมาเยี่ยมและยินดีคุย ใช้แบบครอบครัว (หน้า 6) · ก่อนกลับขอบคุณและขอช่องทางติดต่อ", "ตามโอกาส"],
    ], [0.7, 13.6, 2.7], size=S)

    head(doc, "ลำดับกับทุกคน")
    rich(doc, [("① ", True, BRAND), ("เซ็นใบยินยอม   ", False, None),
               ("② ", True, BRAND), ("ถามปัญหา ยังไม่พูดถึงโปรแกรม   ", False, None),
               ("③ ", True, BRAND), ("อ่านประโยคแนะนำ ให้เล่ากลับ   ", False, None),
               ("④ ", True, BRAND), ("บัตรฟีเจอร์ 3 ใบ   ", False, None),
               ("⑤ ", True, BRAND), ("ขอติดต่อกลับ", False, None)], size=S, after=2)
    box(doc, "ประโยคเปิด: \"กำลังศึกษาเรื่องการใช้ชีวิตของผู้สูงอายุ ไม่มีคำตอบถูกผิด\"  ·  แล้วหลังจบตอนที่ 1 อ่านประโยคแนะนำ",
        SCRIPT_30S)
    box(doc, "ห้าม",
        "ให้ผู้สูงอายุทดสอบร่างกายจริง · ถ่ายเห็นใบหน้า · กรอกแทนคนตอบ · เรียกว่ากล้องจับการล้ม", fill="FEF2F2", color=STOP)

    head(doc, "แผนสำรอง")
    table(doc, ["ถ้าเจอแบบนี้", "ให้ทำ"], [
        ["ไม่อนุญาตให้คุยกับผู้พักอาศัย", "คุยเจ้าหน้าที่ 2 คน แล้วหาผู้สูงอายุจากคนรู้จักหรือเพื่อนบ้านวันเดียวกัน"],
        ["ยังไม่อนุญาตเลย", "ยื่นหนังสือไว้ เก็บจากคนรู้จัก: ผู้สูงอายุ 3 · ลูกหลาน 2–3 · อสม./พยาบาล"],
        ["ได้ไม่ครบ 5 คน", "เขียนจำนวนตามจริง ไปเติมกลุ่มครอบครัววันที่ 16"],
    ], [5.2, 11.8], size=S)

    two_col(doc, evening, deadline)
    tighten(doc)

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    out = build()
    print("เขียน", os.path.relpath(out, ROOT))
    if "--no-pdf" not in sys.argv:
        pages = kit.to_pdf(OUT_DOCX, OUT_PDF)
        print("เขียน", os.path.relpath(OUT_PDF, ROOT), "·", pages, "หน้า")
