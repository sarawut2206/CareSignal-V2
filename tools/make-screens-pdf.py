# -*- coding: utf-8 -*-
"""
make-screens-pdf.py — ภาพหน้าจอ CareSignal พร้อมคำบรรยาย ไว้เปิดให้ผู้ถูกสัมภาษณ์ดู
-----------------------------------------------------------------------------------
python tools/make-screens-pdf.py <โฟลเดอร์ภาพ>
  → validation/ภาพหน้าจอ-CareSignal.docx และ .pdf  (หน้าละหนึ่งจอ)

ภาพมาจากการแคปหน้าจอจริงของไฟล์ในโปรเจกต์นี้ด้วย Chrome headless
หน้าที่มีแถบ "ข้อมูลสาธิต" คือข้อมูลสังเคราะห์ ไม่ใช่ผู้ป่วยจริง — คำบรรยายบอกไว้ทุกหน้า
"""
import os, sys, importlib.util
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("kit", os.path.join(HERE, "make-bangkhae-kit.py"))
kit = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(kit)
for _k in ("para", "box", "h1", "set_font", "BRAND", "INK2", "WARN", "FONT", "ROOT", "OUT_DIR"):
    globals()[_k] = getattr(kit, _k)

OUT_DOCX = os.path.join(OUT_DIR, "ภาพหน้าจอ-CareSignal.docx")
OUT_PDF = os.path.join(OUT_DIR, "ภาพหน้าจอ-CareSignal.pdf")

# (ไฟล์ภาพ, หัวเรื่อง, คำบรรยาย, กว้างเป็นเซนติเมตร, เป็นข้อมูลสาธิตไหม)
PAGES = [
    ("02-app-welcome.png", "แอปในมือถือของผู้สูงอายุ",
     "เปิดแอปมาเจอหน้านี้ก่อน ให้เลือกว่าเป็นผู้สูงอายุเจ้าของผลตรวจ หรือเป็นลูกหลานที่ดูแล "
     "ปุ่มล่างมีสี่อย่างเท่านั้น คือ หน้าแรก ตรวจร่างกาย แนวโน้ม และครอบครัว", 10.6, False),
    ("05-app-intro.png", "ก่อนเริ่มวัด แอปบอกให้เตรียมอะไรบ้าง",
     "บอกชัดว่าใช้เก้าอี้แบบไหน ต้องมีที่ว่างเท่าไร ใส่รองเท้าหรือไม่ และบอกเหตุผลว่าทำไมต้องใช้ระยะ 3 เมตร "
     "ถ้าพื้นที่ไม่พอ ข้ามข้อนั้นได้ คะแนนส่วนอื่นยังใช้ได้", 10.6, False),
    ("06-app-safety.png", "ถามความปลอดภัยก่อนทุกครั้ง",
     "ถ้าวันนี้เวียนหัว เจ็บหน้าอก หรือหายใจลำบาก ระบบจะหยุดให้เอง ไม่ให้ทดสอบต่อ "
     "และแนะนำสิ่งที่ควรทำแทน — ความปลอดภัยมาก่อนการเก็บข้อมูล", 10.6, False),
    ("09-vision-assess.png", "การประเมินแบ่งเป็นข้อ ๆ ทำทีละข้อ",
     "มี 7 หัวข้อ เช่น ประวัติหกล้ม ยาที่ใช้ประจำ ลุกเดิน 3 เมตร หยุดกลางทางแล้วกลับมาทำต่อได้ "
     "ผลจะถูกส่งเข้าระบบเมื่อกดบันทึกเท่านั้น", 15.5, False),
    ("10-staff.png", "เมื่อระบบพบสัญญาณเสี่ยง ใครทำอะไรต่อ",
     "หน้าจอของเจ้าหน้าที่ดูแล เห็นว่าใครต้องติดต่อด่วน ใครรอผลจากแพทย์ และงานไหนเกินกำหนด "
     "การตัดสินใจส่งต่อทุกครั้งเป็นของคน ไม่ใช่ของโปรแกรม", 16.5, True),
    ("11-dashboard.png", "ภาพรวมที่บริษัทประกันเห็น",
     "เห็นเป็นภาพรวมของกลุ่มว่ามีกี่คนเสี่ยงระดับไหน และงานค้างอยู่ตรงไหน "
     "หน้านี้ไม่มีการคำนวณเบี้ย ไม่มีการพิจารณารับประกัน และไม่มีการตัดสินสินไหม", 16.5, True),
    ("12-journey.png", "ตัวอย่างเส้นทางหนึ่งเคส ตั้งแต่ลูกสาวแจ้งเหตุ",
     "ไล่ให้เห็นทีละวันว่าเกิดอะไรขึ้นบ้าง ตั้งแต่ลูกสาวแจ้งว่าแม่ล้มในห้องน้ำ ไปจนถึงการประเมินซ้ำที่บ้าน "
     "และการส่งต่อเภสัชกร นักกายภาพบำบัด แพทย์ และพยาบาล", 16.5, True),
    ("01-overview.png", "หน้าเว็บรวมของโครงการ",
     "หน้านี้อธิบายภาพใหญ่ของระบบ และเป็นทางเข้าไปยังหน้าจอของแต่ละบทบาท", 16.5, False),
]


def build(shots):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.2)
    sec.footer_distance = Cm(0.6)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run("ภาพหน้าจอ CareSignal · ต้นแบบที่ส่งประกวด OIC InsurTech Award 2026 · หน้า "), 10.5, False, INK2)
    r = fp.add_run(); set_font(r, 10.5, False, INK2)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._element.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._element.append(it)

    first = True
    for fn, title, caption, width, demo in PAGES:
        path = os.path.join(shots, fn)
        if not os.path.exists(path):
            print("ข้าม (ไม่พบภาพ)", fn); continue
        h1(doc, title, None, first=first); first = False
        para(doc, caption, size=14, after=5)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        p.add_run().add_picture(path, width=Cm(width))
        if demo:
            para(doc, "แถบสีน้ำตาลด้านบนของภาพคือป้ายข้อมูลสาธิต ทุกชื่อในภาพเป็นข้อมูลสังเคราะห์ ไม่ใช่คนจริง",
                 size=12.5, color=WARN, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


if __name__ == "__main__":
    shots = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "_shots")
    out = build(shots)
    print("เขียน", os.path.relpath(out, ROOT))
    if "--no-pdf" not in sys.argv:
        print("เขียน", os.path.relpath(OUT_PDF, ROOT), "·", kit.to_pdf(OUT_DOCX, OUT_PDF), "หน้า")
