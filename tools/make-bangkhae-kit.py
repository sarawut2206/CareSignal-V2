# -*- coding: utf-8 -*-
"""
make-bangkhae-kit.py — ชุดสัมภาษณ์ภารกิจที่ 3 ฉบับใช้ที่บ้านบางแค เป็น .docx (+ .pdf ผ่าน Word)
------------------------------------------------------------------------------------------
python tools/make-bangkhae-kit.py            → validation/ชุดสัมภาษณ์-บ้านบางแค.docx และ .pdf

เนื้อหาชุดเดียวกับ tools/make-validation-kit.mjs (ฉบับ PDF 10 หน้า) แต่เป็นไฟล์ Word ที่แก้ได้
และเพิ่มหน้าแรกสำหรับการลงพื้นที่บ้านบางแค:
  · ใครต้องสัมภาษณ์ กี่คน ใช้แบบไหน และต้องพิมพ์อะไรกี่แผ่น
  · ประโยคแนะนำ CareSignal 30 วินาทีที่อธิบายระบบถูกต้อง (ประเมินความเสี่ยงเป็นรอบ ไม่ใช่กล้องจับการล้ม)
  · ข้อควรระวังเฉพาะบ้านพักผู้สูงอายุ (ความยินยอม ห้ามทดสอบร่างกาย ภาพไม่เห็นหน้า)
  · คำถามเพิ่มสำหรับเจ้าหน้าที่ของศูนย์ฯ

ไม่ต้องใช้ข้อมูลของผู้ให้ข้อมูลคนใดเลย — ทุกช่องเป็นช่องว่างให้กรอกด้วยมือในวันสัมภาษณ์
"""
import os, sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

try:
    from pythainlp.tokenize import word_tokenize
except Exception:
    word_tokenize = None

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "validation")
OUT_DOCX = os.path.join(OUT_DIR, "ชุดสัมภาษณ์-บ้านบางแค.docx")
OUT_PDF = os.path.join(OUT_DIR, "ชุดสัมภาษณ์-บ้านบางแค.pdf")
FONT = "TH Sarabun New"
BRAND = RGBColor(0x14, 0x3A, 0x74)
INK2 = RGBColor(0x4F, 0x5F, 0x78)
WARN = RGBColor(0x92, 0x40, 0x0E)
STOP = RGBColor(0x99, 0x1B, 0x1B)
ZW = "​"
KEEP = ["CareSignal", "บ้านบางแค", "ผู้สูงอายุ", "นักกายภาพบำบัด", "เภสัชกร", "results.json"]

# บัตรฟีเจอร์ — ต้องตรงกับ FEATURES ใน tools/make-validation-kit.mjs ทุกตัวอักษร
FEATURES = [
    ("F1", "ตรวจร่างกายเองที่บ้านด้วยกล้องมือถือ ไม่ต้องเดินทางไปโรงพยาบาล"),
    ("F2", "รู้ผลทันทีว่าเสี่ยงล้มระดับไหน พร้อมบอกเหตุผลว่าทำไม"),
    ("F3", "ลูกหลานได้รับแจ้งเตือนเมื่อผลแย่ลง โดยที่เราเป็นคนอนุญาต"),
    ("F4", "มีเจ้าหน้าที่โทรกลับภายใน 1–3 วันเมื่อระบบพบสัญญาณเสี่ยง"),
    ("F5", "เภสัชกรช่วยดูว่ายาที่กินอยู่ตัวไหนทำให้เวียนหัวหรือเสี่ยงล้ม"),
    ("F6", "ส่งต่อนักกายภาพบำบัดพร้อมข้อมูล ไม่ต้องเล่าเรื่องใหม่ตั้งแต่ต้น"),
    ("F7", "ดูย้อนหลังได้ว่าเดือนนี้ดีขึ้นหรือแย่ลงกว่าครั้งก่อน"),
    ("F8", "ภาพจากกล้องไม่ออกจากเครื่อง และลบข้อมูลตัวเองได้ทุกเมื่อ"),
]

SCRIPT_30S = ("เป็นโปรแกรมในมือถือ ให้ผู้สูงอายุทำแบบทดสอบง่าย ๆ ประมาณ 10 นาที ทุก 1–3 เดือน "
              "เช่น ลุกนั่งจากเก้าอี้ ลุกเดิน และยืนทรงตัว เพื่อดูว่าเสี่ยงล้มมากขึ้นไหม ก่อนที่จะล้มจริง "
              "ถ้าพบว่าเสี่ยง จะมีเจ้าหน้าที่โทรหา แล้วส่งต่อเภสัชกร นักกายภาพบำบัด หรือแพทย์ตามปัญหา "
              "ครอบครัวเห็นผลที่ผู้เชี่ยวชาญยืนยันแล้ว กล้องไม่เก็บภาพและไม่ส่งภาพออกจากเครื่อง")


# ------------------------------------------------------------ ตัวช่วยจัดรูปแบบ
def thai(text):
    """ฝังอักขระกว้างศูนย์ตามขอบเขตคำ ให้ Word ตัดบรรทัดภาษาไทยได้แม้เครื่องไม่ได้เปิดภาษาไทย"""
    if not word_tokenize or not re.search(r"[฀-๿]", text):
        return text
    out = []
    for seg in re.split(r"(\s+)", text):
        if not seg or seg.isspace() or any(k in seg for k in KEEP) or not re.search(r"[฀-๿]", seg):
            out.append(seg)
            continue
        out.append(ZW.join(word_tokenize(seg, engine="newmm", keep_whitespace=False)))
    return "".join(out)


def set_font(run, size=14, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.append(rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), FONT)
    szcs = OxmlElement("w:szCs"); szcs.set(qn("w:val"), str(int(size * 2))); rpr.append(szcs)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    lang = OxmlElement("w:lang"); lang.set(qn("w:bidi"), "th-TH"); rpr.append(lang)
    if color is not None:
        run.font.color.rgb = color


def para(doc_or_cell, text="", size=14, bold=False, color=None, align=None, before=0, after=2, indent=None, keep=False):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.0
    if indent is not None:
        pf.left_indent = Cm(indent)
    if keep:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    if text:
        set_font(p.add_run(thai(text)), size, bold, color)
    return p


def rich(doc_or_cell, parts, size=14, after=2, indent=None):
    """parts = [(ข้อความ, ตัวหนา, สี)] ในย่อหน้าเดียว"""
    p = doc_or_cell.add_paragraph()
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.0
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    for t, b, c in parts:
        set_font(p.add_run(thai(t)), size, b, c)
    return p


def shade(cell, hex_fill):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def borders(table, color="9AA8BD", size=6, dash=False):
    tbl = table._element
    tblpr = tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "dashed" if dash else "single"); e.set(qn("w:sz"), str(size)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def cell_text(cell, text, size=13, bold=False, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    if align:
        p.alignment = align
    set_font(p.add_run(thai(text)), size, bold, color)
    return p


def table(doc, header, rows, widths, size=12.5, head_fill="DCEBFB", row_height=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    borders(t)
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.width = Cm(widths[i]); cell_text(c, h, size, True, BRAND); shade(c, head_fill)
    for r, row in enumerate(rows, start=1):
        for i, v in enumerate(row):
            c = t.rows[r].cells[i]; c.width = Cm(widths[i]); cell_text(c, v, size, i == 0 and bool(v) and len(v) > 3)
        if row_height:
            tr = t.rows[r]._tr; trpr = tr.get_or_add_trPr()
            h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(row_height * 567))); h.set(qn("w:hRule"), "atLeast"); trpr.append(h)
    para(doc, "", size=3, after=0)
    return t


def box(doc, title, body, fill="EAF3FC", color=BRAND):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, color="C9DAF3" if fill == "EAF3FC" else "E5C07B" if fill == "FFF7ED" else "F5A3A3")
    c = t.rows[0].cells[0]; c.width = Cm(17); shade(c, fill)
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    set_font(p.add_run(thai(title)), 14, True, color)
    for line in (body if isinstance(body, list) else [body]):
        q = c.add_paragraph(); q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = 1.0
        set_font(q.add_run(thai(line)), 13)
    para(doc, "", size=3, after=0)


def h1(doc, title, sub=None, page_label=None, first=False):
    p = para(doc, title, size=20, bold=True, color=BRAND, after=0, keep=True)
    if not first:
        p.paragraph_format.page_break_before = True   # ขึ้นหน้าใหม่ที่หัวเรื่อง ไม่ใช้ย่อหน้าว่าง ซึ่งกินหนึ่งบรรทัดบนสุดของหน้า
    if sub:
        para(doc, sub, size=13.5, color=INK2, after=0, keep=True)
    rule = doc.add_paragraph(); rule.paragraph_format.space_after = Pt(3)
    set_font(rule.add_run(""), 4)
    ppr = rule._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "12"); bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "1D4E9A")
    pb.append(bt); ppr.append(pb)
    return p


def h2(doc, text):
    return para(doc, text, size=15, bold=True, color=BRAND, before=4, after=1, keep=True)


def lines(doc, n=2, indent=0.6):
    for _ in range(n):
        para(doc, "_" * 100, size=11, color=RGBColor(0xB8, 0xC4, 0xD4), after=0, indent=indent)


def q(doc, no, question, hint=None, n=2):
    rich(doc, [(f"{no}. ", True, BRAND), (question, True, None)], size=13.5, after=0)
    if hint:
        para(doc, hint, size=12, color=INK2, after=0, indent=0.6)
    lines(doc, n)
    para(doc, "", size=3, after=0)


def ticks(doc, items, label=None):
    parts = []
    if label:
        parts.append((label + "   ", True, None))
    for it in items:
        parts.append(("☐ " + it + "      ", False, None))
    rich(doc, parts, size=14, after=1)


def fields(doc, labels):
    rich(doc, [(f"{l} ........................   ", False, None) for l in labels], size=14, after=3)


# ------------------------------------------------------------ หน้าต่าง ๆ
def page_field_day(doc):
    h1(doc, "ชุดสัมภาษณ์ CareSignal · ฉบับลงพื้นที่บ้านบางแค",
       "ภารกิจที่ 3 Idea Validation · OIC InsurTech Award 2026 · ส่งภายใน 16 กันยายน 2569", first=True)
    box(doc, "กติกาข้อเดียวที่ห้ามผิด",
        "ข้อมูลทุกตัวในใบสรุปต้องมาจากคนจริงที่ตอบจริง ห้ามกรอกแทน ห้ามเดา ห้ามเติมให้ครบ · "
        "ถ้าได้ 3 คนก็เขียน 3 คน กรรมการให้คะแนนความจริงของข้อมูล ไม่ได้ให้คะแนนจำนวน")

    h2(doc, "สัมภาษณ์ใคร กี่คน ใช้แบบไหน")
    table(doc, ["กลุ่ม", "ใคร", "เป้าหมาย", "ใช้แบบ", "เก็บที่ไหน"], [
        ["ก", "ผู้สูงอายุที่สื่อสารและให้ความยินยอมได้เอง", "3 คน", "หน้า 4–5", "บ้านบางแค"],
        ["ค", "พยาบาล ผู้ช่วยพยาบาล นักกายภาพบำบัด นักสังคมสงเคราะห์", "2 คน", "หน้า 7", "บ้านบางแค"],
        ["ข", "ญาติที่มาเยี่ยม หรือผู้ปกครองนักเรียนที่ดูแลพ่อแม่อยู่บ้าน", "1–3 คน", "หน้า 6 หรือ 9", "บางแค หรือโทร/LINE ภายใน 16 ก.ย."],
    ], [1.3, 7.2, 2.0, 2.6, 3.9])

    h2(doc, "ต้องพิมพ์อะไร กี่แผ่น")
    table(doc, ["หน้า", "เอกสาร", "จำนวน"], [
        ["3", "ใบยินยอมให้สัมภาษณ์", "6 แผ่น (คนละแผ่น)"],
        ["4–5", "แบบสัมภาษณ์ ผู้สูงอายุ", "4 ชุด (เผื่อ 1)"],
        ["6", "แบบสัมภาษณ์ ครอบครัว/ผู้ดูแล", "2 แผ่น"],
        ["7", "แบบสัมภาษณ์ เจ้าหน้าที่ (มีคำถามเพิ่มสำหรับศูนย์ฯ)", "3 แผ่น"],
        ["8", "บัตรฟีเจอร์ 8 ใบ (ตัดตามเส้นประ)", "1 แผ่น ใช้ซ้ำทุกคน"],
        ["9", "แบบสอบถามตอบเอง (สำหรับครอบครัวทาง LINE)", "ส่งเป็นไฟล์ได้"],
        ["10", "แบบสรุปรายคน กรอกทันทีหลังคุยจบ", "6 แผ่น"],
    ], [1.6, 11.0, 4.4])

    h2(doc, "ประโยคแนะนำ CareSignal (อ่านหลังจบตอนที่ 1 เท่านั้น)")
    box(doc, "อ่านตามนี้ ใช้เวลาประมาณ 30 วินาที", SCRIPT_30S)
    box(doc, "ห้ามอธิบายว่าเป็นกล้องจับการล้มหรือระบบแจ้งเตือนเมื่อล้ม",
        "ระบบของเราประเมินความเสี่ยงเป็นรอบ ๆ ก่อนเกิดเหตุ ถ้าเล่าผิด คำตอบที่ได้จะนำขึ้นเวทีไม่ได้", fill="FFF7ED", color=WARN)

    h2(doc, "ข้อควรระวังเฉพาะบ้านพักผู้สูงอายุ")
    for t in [
        "ยื่นหนังสือขออนุญาตกับเจ้าหน้าที่ของศูนย์ฯ ก่อนเริ่ม ถามกติกาการถ่ายภาพ (ถ่ายเฉพาะมือหรือด้านหลัง ไม่เห็นใบหน้า)",
        "ให้เจ้าหน้าที่ช่วยแนะนำผู้สูงอายุที่สื่อสารได้ชัด ถ้าไม่แน่ใจว่าท่านเข้าใจใบยินยอม ให้งดสัมภาษณ์ท่านนั้น",
        "ไม่ให้ผู้สูงอายุทดสอบลุกนั่ง เดิน หรือทรงตัวจริง ต้องการความคิดเห็นเท่านั้น ถ้าจะให้เห็นภาพ ให้สมาชิกทีมสาธิตเอง",
        "คุยคนละ 15–20 นาที ท่านเหนื่อยให้หยุด · ในใบสรุปเขียนตรง ๆ ว่าผู้สูงอายุในบ้านพักไม่ใช่กลุ่มหลัก (กลุ่มหลักอยู่บ้านตัวเอง)",
    ]:
        rich(doc, [("•  ", True, BRAND), (t, False, None)], size=13.5, after=0, indent=0.3)


def page_plan(doc):
    h1(doc, "สิ่งที่ต้องส่ง และวิธีถามให้ได้ความจริง", "อ่านก่อนเริ่มสัมภาษณ์คนแรก")
    h2(doc, "สิ่งที่ต้องส่ง 5 หัวข้อ ได้มาจากคำถามข้อไหน")
    table(doc, ["หัวข้อที่ต้องมีในใบสรุป", "ได้จากตรงไหนในชุดนี้"], [
        ["กลุ่มเป้าหมายและจำนวนผู้ให้ข้อมูล", "ใบยินยอมหน้า 3 นับจำนวน และแบ่งกลุ่ม ก ข ค"],
        ["ปัญหาที่ทีมแก้ เป็นปัญหาจริงหรือไม่", "ตอนที่ 1 ของทุกแบบสัมภาษณ์ (ถามก่อนบอกไอเดีย) นับว่ากี่คนเคยล้มหรือเกือบล้ม และกี่คนไม่เคยถูกประเมิน"],
        ["Solution ช่วยแก้ปัญหาได้หรือไม่", "ตอนที่ 2 ข้อ 8 (อธิบายได้เองไหมว่าโปรแกรมทำอะไร) และตอนที่ 4 (ให้ช่องทางติดต่อกลับหรือไม่)"],
        ["ฟีเจอร์ใดที่ผู้ใช้ให้ความสำคัญที่สุด", "บัตรฟีเจอร์หน้า 8 เลือก 3 ใบ เรียง 1-2-3 แปลงเป็นคะแนน 3/2/1 แล้วรวมทุกคน"],
        ["จะนำ Feedback ไปใช้อย่างไร", "ช่อง สิ่งที่เราจะเปลี่ยน ในแบบสรุปรายคนหน้า 10"],
    ], [6.0, 11.0])
    h2(doc, "ห้าข้อห้าม")
    for t in [
        "อย่าอธิบายไอเดียก่อนถามปัญหา ถ้าเล่าก่อน คำตอบทุกข้อหลังจากนั้นจะเอียงตามสิ่งที่เราเล่า",
        "อย่าถามว่า ไอเดียนี้ดีไหม จะได้แต่คำชมที่เอาไปใช้ไม่ได้",
        "อย่าถามว่า ท่านจะใช้ไหม คนตอบว่าใช้แล้วไม่ได้ใช้เป็นเรื่องปกติมาก ให้ขอสิ่งที่เป็นการกระทำแทน เช่น ขอเบอร์ติดต่อกลับ",
        "อย่าถามคำถามชี้นำ เช่น ท่านคงกังวลเรื่องหกล้มใช่ไหม ให้ถามว่า มีเรื่องอะไรที่ทำให้ต้องระวังเป็นพิเศษบ้าง",
        "อย่าเถียงหรืออธิบายเพิ่มเมื่อเขาบอกว่าไม่มีปัญหา ให้จดไว้ตรง ๆ นั่นคือข้อมูลที่มีค่าที่สุด",
    ]:
        rich(doc, [("✗  ", True, STOP), (t, False, None)], size=13.5, after=0, indent=0.3)
    h2(doc, "ห้าข้อควรทำ")
    for t in [
        "ถามถึงเหตุการณ์ที่เกิดขึ้นจริงในอดีต ไม่ใช่สิ่งที่เขาคิดว่าจะทำในอนาคต",
        "เงียบให้เป็น ถามแล้วรออย่างน้อย 5 วินาที คำตอบที่ดีที่สุดมักตามหลังความเงียบ",
        "จดคำพูดจริงเป็นประโยค ใส่เครื่องหมายคำพูดไว้ อย่าสรุปเป็นความคิดของเราเอง",
        "ถ้าคำตอบคลุมเครือ ให้ถามต่อว่า เล่าให้ฟังหน่อยว่าครั้งล่าสุดเป็นอย่างไร",
        "สัมภาษณ์ทีละคน อย่าถามพร้อมกันเป็นกลุ่มเพราะจะตอบตามกัน",
    ]:
        rich(doc, [("✓  ", True, RGBColor(0x16, 0x65, 0x34)), (t, False, None)], size=13.5, after=0, indent=0.3)
    h2(doc, "แบ่งเวลา (ผู้สูงอายุ 15–20 นาที · เจ้าหน้าที่ 20–30 นาที)")
    table(doc, ["ช่วง", "ตอนที่", "ห้ามลืม"], [
        ["เริ่ม", "แนะนำตัว", "ขอความยินยอม ขอถ่ายภาพ"],
        ["ส่วนแรก", "ตอนที่ 1", "ถามชีวิตและเรื่องที่เกิดจริง ยังไม่พูดถึงระบบเรา"],
        ["ส่วนที่สอง", "ตอนที่ 2", "อ่านประโยคแนะนำ 30 วินาที แล้วให้เขาอธิบายกลับ"],
        ["ส่วนที่สาม", "ตอนที่ 3", "บัตรฟีเจอร์ เลือก 3 ใบ เรียง 1-2-3"],
        ["ท้าย", "ตอนที่ 4", "ขอสิ่งที่เป็นการกระทำ แล้วขอบคุณ"],
    ], [3.0, 3.0, 11.0])
    box(doc, "สิ่งที่ทำให้ข้อมูลใช้ไม่ได้เลย",
        "กรอกแบบสอบถามแทนผู้ให้ข้อมูล · ถามคนในทีมเองแล้วนับเป็นผู้ใช้ · เอาคำตอบของคนหนึ่งไปเติมให้อีกคน · "
        "ถ่ายภาพโดยไม่ขออนุญาต · เขียนในใบสรุปว่าถามมากกว่าที่ถามจริง", fill="FEF2F2", color=STOP)


def page_consent(doc):
    h1(doc, "ใบยินยอมให้สัมภาษณ์ บันทึก และถ่ายภาพ", "หนึ่งใบต่อผู้ให้ข้อมูลหนึ่งคน · ถ้าอ่านเองไม่สะดวก ให้ผู้สัมภาษณ์อ่านให้ฟังทั้งหมดก่อนลงชื่อ")
    para(doc, "ทีม CareSignal กำลังพัฒนาโปรแกรมช่วยเฝ้าระวังความเสี่ยงหกล้มในผู้สูงอายุ และขอสอบถามความคิดเห็นของท่านเพื่อนำไปปรับปรุง "
              "การสัมภาษณ์ใช้เวลาประมาณ 15–30 นาที เป็นการพูดคุยเท่านั้น ไม่มีการทดสอบร่างกาย", size=14, after=4)
    table(doc, ["ข้อมูลที่เก็บ", "ข้อมูลที่ไม่เก็บ"], [
        ["รหัสผู้ให้ข้อมูล เช่น R01 (ไม่ใช่ชื่อจริง)", "ชื่อ นามสกุล เลขประจำตัวประชาชน ที่อยู่"],
        ["ช่วงอายุ เพศ และบทบาท (ผู้สูงอายุ ครอบครัว เจ้าหน้าที่)", "ประวัติการรักษาหรือผลตรวจสุขภาพ"],
        ["ความคิดเห็นและคำพูดของท่านต่อโปรแกรม", "ข้อมูลกรมธรรม์ประกันภัยของท่าน"],
        ["ภาพบรรยากาศการสัมภาษณ์ ถ้าท่านอนุญาต", "เบอร์โทร ยกเว้นท่านให้ไว้เองเพื่อให้ติดต่อกลับ"],
    ], [8.5, 8.5])
    h2(doc, "สิทธิของท่าน")
    for t in [
        "ไม่ตอบข้อไหนก็ได้ และหยุดกลางคันได้ทุกเมื่อ",
        "ขอให้ลบข้อมูลและภาพของท่านได้ภายหลัง โดยแจ้งรหัสผู้ให้ข้อมูลกับทีมงาน",
        "การเข้าร่วมหรือไม่เข้าร่วม ไม่มีผลต่อสิทธิ์การรักษา สวัสดิการ หรือกรมธรรม์ใด ๆ ของท่าน",
        "ข้อมูลนี้ใช้เพื่อพัฒนาโปรแกรมและส่งเป็นผลงานในการประกวดเท่านั้น ไม่ขายและไม่ส่งต่อให้บริษัทประกัน",
    ]:
        rich(doc, [("•  ", True, BRAND), (t, False, None)], size=13.5, after=0, indent=0.3)
    h2(doc, "คำยินยอม ให้ผู้ให้ข้อมูลติ๊กเอง")
    for t in ["ยินยอมให้สัมภาษณ์และจดบันทึก", "ยินยอมให้บันทึกเสียงระหว่างสัมภาษณ์",
              "ยินยอมให้ถ่ายภาพ และนำภาพไปใช้ประกอบผลงานที่ส่งประกวด",
              "ยินยอมให้ถ่ายภาพ แต่ขอให้เห็นเฉพาะมือหรือด้านหลัง ไม่เห็นใบหน้า"]:
        ticks(doc, [t])
    para(doc, "", size=6)
    t = doc.add_table(rows=2, cols=2); borders(t)
    for (r, c, label) in [(0, 0, "รหัสผู้ให้ข้อมูล"), (0, 1, "วันที่"), (1, 0, "ลงชื่อผู้ให้ข้อมูล"), (1, 1, "ลงชื่อผู้สัมภาษณ์")]:
        cell = t.rows[r].cells[c]; cell.width = Cm(8.5)
        cell_text(cell, label, 13, True, INK2)
        pp = cell.add_paragraph(); pp.paragraph_format.space_after = Pt(0)
        set_font(pp.add_run("\n.............................................................."), 14)
    para(doc, "", size=6)
    fields(doc, ["ขอถอนตัวหรือขอลบข้อมูล เมื่อวันที่", "ผู้รับแจ้ง"])


def header_row(doc, extra=None):
    labels = ["รหัส", "อายุ ....... ปี", "วันที่", "ผู้สัมภาษณ์"]
    rich(doc, [(l + (" ................   " if "ปี" not in l else "   "), False, None) for l in labels], size=14, after=4)


def page_elder_a(doc):
    h1(doc, "แบบสัมภาษณ์ กลุ่ม ก · ผู้สูงอายุ", "ตอนที่ 1 ถามก่อนบอกไอเดีย — ห้ามพูดถึงโปรแกรมของเราในหน้านี้")
    header_row(doc)
    box(doc, "หน้านี้ห้ามพูดถึง CareSignal",
        "เปิดด้วยประโยคนี้: ผมกำลังศึกษาเรื่องการใช้ชีวิตของผู้สูงอายุ ขอถามเรื่องชีวิตประจำวันของคุณลุงคุณป้าสักครู่ ไม่มีคำตอบถูกผิดครับ",
        fill="FFF7ED", color=WARN)
    q(doc, 1, "ตอนนี้คุณลุงคุณป้าอยู่ที่นี่มานานเท่าไร แล้ววันธรรมดาวันหนึ่งเป็นอย่างไร",
      "ให้เล่าสบาย ๆ ก่อน จดว่าเดินเองได้ไหม ใช้ไม้เท้าหรือไม่ และทำกิจกรรมอะไรบ้าง", 2)
    q(doc, 2, "ในหนึ่งปีที่ผ่านมา มีครั้งไหนที่ล้มหรือเกือบล้มบ้างไหม เล่าให้ฟังหน่อยว่าวันนั้นเกิดอะไรขึ้น",
      "ถ้ามี ถามต่อ: ตอนนั้นอยู่ตรงไหน กำลังทำอะไร หลังจากนั้นบอกใคร มีใครพาไปหาหมอไหม · ถ้าไม่มี ข้ามไปข้อ 3 อย่าคะยั้นคะยอ", 3)
    q(doc, 3, "ทุกวันนี้มีอะไรที่ทำได้ยากขึ้นกว่าเมื่อสองปีก่อนไหม",
      "เช่น ลุกจากเก้าอี้ ขึ้นบันได เข้าห้องน้ำตอนกลางคืน หิ้วของ", 2)
    q(doc, 4, "ครั้งสุดท้ายที่ตรวจสุขภาพ มีใครวัดเรื่องการทรงตัวหรือกำลังขาบ้างไหม",
      "ถ้าตอบว่าไม่มี ให้จดว่า ไม่เคยถูกประเมิน — ตัวเลขนี้ใช้ตอบว่าปัญหานี้ยังไม่มีใครดูแลจริงหรือไม่", 2)
    q(doc, 5, "ถ้าอยากรู้ว่าตัวเองเสี่ยงล้มแค่ไหน ตอนนี้ทำอย่างไร",
      "คำตอบว่า ไม่รู้เหมือนกัน หรือ ไม่เคยคิด ก็เป็นคำตอบที่ใช้ได้ ให้จดตามที่เขาพูด", 2)
    q(doc, 6, "ในมือถือมีแอปอะไรที่ใช้เองบ้าง แล้วใครเป็นคนลงให้",
      "บอกเราว่าใครควรเป็นคนตั้งค่าครั้งแรก ถ้าไม่มีมือถือให้จดไว้ด้วย", 2)
    q(doc, 7, "เรื่องสุขภาพของคุณลุงคุณป้า ใครเป็นคนที่รู้มากที่สุดรองจากตัวเอง", None, 2)


def page_elder_b(doc):
    h1(doc, "แบบสัมภาษณ์ กลุ่ม ก · ผู้สูงอายุ (ต่อ)", "ตอนที่ 2 ถึง 4 · อ่านประโยคแนะนำ 30 วินาทีจากหน้า 1 ก่อนถาม")
    rich(doc, [("รหัส ................      วันที่ ................", False, None)], size=14, after=4)
    h2(doc, "ตอนที่ 2 · หลังอ่านประโยคแนะนำ (หรือเปิดวิดีโอ)")
    q(doc, 8, "จากที่ฟัง คุณลุงคุณป้าคิดว่าโปรแกรมนี้ทำอะไร ลองเล่าให้ฟังหน่อย",
      "อย่าช่วยเฉลย · อธิบายได้ใกล้เคียง จด เข้าใจ · อธิบายผิดหรือเงียบ จด ไม่เข้าใจ แล้วจดคำที่เขาใช้ไว้ด้วย", 2)
    q(doc, 9, "ตรงไหนที่ฟังแล้วยังไม่เข้าใจ หรือรู้สึกว่ายาก", None, 2)
    q(doc, 10, "ถ้าจะทำแบบนี้จริง อะไรคืออุปสรรค", "เช่น ไม่มีที่วางมือถือ ไม่มีคนช่วย กลัวทำพัง ห้องแคบ", 2)
    h2(doc, "ตอนที่ 3 · บัตรฟีเจอร์ วางบัตร 8 ใบให้เลือก 3 ใบ")
    para(doc, "วางบัตรทั้ง 8 ใบให้เห็นพร้อมกัน อ่านให้ฟังทีละใบถ้าจำเป็น แล้วให้เลือก 3 ใบที่มีประโยชน์กับตัวเองมากที่สุด เรียงว่าใบไหนสำคัญที่สุด", size=13.5, after=3)
    table(doc, ["อันดับ", "รหัสบัตร", "ถามว่า ทำไมถึงเลือกใบนี้"], [["1", "", ""], ["2", "", ""], ["3", "", ""]], [2.0, 2.6, 12.4], row_height=0.9)
    q(doc, 11, "มีใบไหนที่คิดว่าไม่จำเป็นเลยไหม เพราะอะไร", None, 2)
    h2(doc, "ตอนที่ 4 · ขอสิ่งที่เป็นการกระทำ ไม่ใช่คำชม")
    ticks(doc, ["ยินดีให้ติดต่อกลับเพื่อมาลองใช้ (ขอช่องทางไว้ในใบยินยอม)"])
    ticks(doc, ["มีลูกหลานหรือเจ้าหน้าที่ที่ช่วยตั้งเครื่องให้ครั้งแรกได้"])
    ticks(doc, ["ถ้าบริษัทประกันแถมมาให้ฟรีกับกรมธรรม์ จะลองใช้"])
    rich(doc, [("ถ้าต้องจ่ายเอง เดือนละเท่าไรถึงจะคุ้ม (ให้เขาบอกตัวเลขเอง อย่าเสนอราคาก่อน) ............ บาท", False, None)], size=14, after=3)
    h2(doc, "คำพูดจริงที่ควรจดไว้ (ใส่เครื่องหมายคำพูด)")
    lines(doc, 3, indent=0)


def page_family(doc):
    h1(doc, "แบบสัมภาษณ์ กลุ่ม ข · ครอบครัวหรือผู้ดูแล", "ตอนที่ 1 ถามก่อนบอกไอเดียเหมือนเดิม · ใช้กับญาติที่มาเยี่ยม หรือโทรสัมภาษณ์")
    header_row(doc)
    q(doc, 1, "ตอนนี้ท่านรู้ได้อย่างไรว่าคุณพ่อคุณแม่เป็นอย่างไรในแต่ละวัน",
      "จดกลไกที่ใช้อยู่จริง เช่น โทรทุกเย็น ให้เพื่อนบ้านช่วยดู กล้องวงจรปิด", 1)
    q(doc, 2, "ครั้งล่าสุดที่ท่านเป็นห่วงเรื่องการล้ม เกิดอะไรขึ้น แล้วท่านทำอะไรต่อ", None, 2)
    q(doc, 3, "เคยมีเรื่องที่ท่านมารู้ทีหลังไหม รู้ช้าไปประมาณกี่วัน", "ถ้าตอบว่ารู้ทันทีเสมอ ให้จดตามนั้น", 1)
    q(doc, 4, "ทุกวันนี้ท่านใช้เวลาหรือเงินไปกับเรื่องนี้อย่างไรบ้าง", "เช่น ลางานพาไปหาหมอ จ้างคนดูแล ค่าเดินทาง", 1)
    q(doc, 5, "ถ้ามีข้อความแจ้งว่า ผลเดือนนี้ของคุณแม่แย่ลง ท่านจะทำอะไรต่อ",
      "ถ้าตอบว่าไม่รู้จะทำอะไร แปลว่าการแจ้งเตือนอย่างเดียวไม่มีค่า ต้องมีคนรับงานต่อ", 1)
    q(doc, 6, "อะไรที่จะทำให้ท่านรำคาญจนปิดการแจ้งเตือนทิ้ง", None, 1)
    h2(doc, "ตอนที่ 2 ถึง 4 · อ่านประโยคแนะนำ แล้วถาม")
    q(doc, 7, "จากที่ฟัง ท่านคิดว่าโปรแกรมนี้ทำอะไร (ให้เขาอธิบายเอง อย่าเฉลย)", None, 1)
    table(doc, ["อันดับบัตรฟีเจอร์", "รหัสบัตร", "เหตุผล"], [["1", "", ""], ["2", "", ""], ["3", "", ""]], [3.6, 2.6, 10.8], row_height=0.8)
    ticks(doc, ["ยินดีให้ติดต่อกลับเพื่อมาลองใช้กับคนในบ้าน      ☐ ยินดีเป็นคนตั้งเครื่องให้ผู้สูงอายุครั้งแรก"])
    h2(doc, "คำพูดจริงที่ควรจดไว้")
    lines(doc, 2, indent=0)


def page_staff(doc):
    h1(doc, "แบบสัมภาษณ์ กลุ่ม ค · เจ้าหน้าที่", "พยาบาล ผู้ช่วยพยาบาล นักกายภาพบำบัด นักสังคมสงเคราะห์ · บอกได้ว่าระบบจะถูกใช้จริงหรือกลายเป็นภาระ")
    header_row(doc)
    rich(doc, [("ตำแหน่ง ..................................................      ดูแลผู้สูงอายุประมาณ ............ คน", False, None)], size=14, after=4)
    q(doc, 1, "ในงานของท่าน ตอนนี้รู้ได้อย่างไรว่าผู้สูงอายุคนไหนเสี่ยงล้ม",
      "จดเครื่องมือที่ใช้อยู่จริง เช่น แบบประเมินกระดาษ Thai-FRAT การสังเกต หรือรอให้เกิดเหตุ · ประเมินบ่อยแค่ไหน", 2)
    q(doc, 2, "เมื่อรู้ว่าใครเสี่ยง หรือเมื่อมีคนล้ม เกิดอะไรขึ้นต่อ ใครรับผิดชอบ และใช้เวลานานแค่ไหน",
      "ถ้าคำตอบคือไม่มีอะไรเกิดขึ้นต่อ ให้จดไว้ตรง ๆ", 2)
    q(doc, 3, "เวลาส่งต่อแพทย์ นักกายภาพ หรือโรงพยาบาล ข้อมูลที่ส่งไปมักขาดอะไร", None, 2)
    q(doc, 4, "อะไรที่ทำให้การติดตามไม่เกิดขึ้นจริง", "เช่น คนไม่พอ ไม่มีข้อมูล ไม่มีใครสั่ง", 1)
    q(doc, 5, "ถ้ามีแบบฟอร์มส่งต่อที่ระบบเติมผลประเมินให้บางส่วนแล้ว จะช่วยประหยัดเวลาหรือเพิ่มงาน",
      "ถามตรง ๆ แบบนี้ เพราะคำตอบว่า เพิ่มงาน มีค่ากับเรามากกว่าคำชม", 2)
    q(doc, 6, "ในแบบฟอร์มนั้นต้องมีข้อมูลอะไรบ้าง ท่านถึงจะทำงานต่อได้เลยโดยไม่ต้องโทรถามซ้ำ", None, 1)
    q(doc, 7, "ถ้าผลจากกล้องกับสิ่งที่ท่านเห็นด้วยตาไม่ตรงกัน ท่านจะเชื่ออะไร และอยากให้ระบบบอกอะไรเพิ่ม", None, 1)
    table(doc, ["อันดับบัตรฟีเจอร์", "รหัสบัตร", "เหตุผล"], [["1", "", ""], ["2", "", ""], ["3", "", ""]], [3.6, 2.6, 10.8], row_height=0.8)
    h2(doc, "คำพูดจริงที่ควรจดไว้")
    lines(doc, 2, indent=0)


def page_cards(doc):
    h1(doc, "บัตรฟีเจอร์ 8 ใบ", "ตัดตามเส้นประ วางให้ผู้ให้ข้อมูลเลือก 3 ใบ แล้วเรียง 1-2-3")
    t = doc.add_table(rows=4, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, color="7B8AA1", size=10, dash=True)
    for i, (fid, nm) in enumerate(FEATURES):
        c = t.rows[i // 2].cells[i % 2]; c.width = Cm(8.5)
        cell_text(c, fid, 26, True, BRAND)
        p = c.add_paragraph(); p.paragraph_format.space_after = Pt(10); p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(thai(nm)), 17)
        tr = t.rows[i // 2]._tr; trpr = tr.get_or_add_trPr()
        if trpr.find(qn("w:trHeight")) is None:
            h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(4.2 * 567))); h.set(qn("w:hRule"), "atLeast"); trpr.append(h)
    para(doc, "", size=8)
    box(doc, "วิธีนับคะแนน",
        "อันดับ 1 ได้ 3 คะแนน อันดับ 2 ได้ 2 คะแนน อันดับ 3 ได้ 1 คะแนน รวมคะแนนของทุกคนแล้วเรียง "
        "จะได้คำตอบของหัวข้อ ฟีเจอร์ใดที่ผู้ใช้ให้ความสำคัญมากที่สุด เป็นตัวเลขที่อธิบายที่มาได้")
    box(doc, "อย่าเผลอทำสองอย่างนี้",
        "อย่าอ่านบัตรด้วยน้ำเสียงเชียร์ใบที่เราอยากให้เลือก · อย่าให้เลือกมากกว่า 3 ใบ เพราะจะไม่รู้ว่าอะไรสำคัญกว่ากัน",
        fill="FFF7ED", color=WARN)


def page_survey(doc):
    h1(doc, "แบบสอบถามตอบเอง", "สำหรับผู้ที่ไม่สะดวกให้สัมภาษณ์ เช่น ครอบครัวทาง LINE · ใช้แทนได้ แต่ได้ข้อมูลตื้นกว่า")
    ticks(doc, ["ผู้สูงอายุ 60 ปีขึ้นไป", "ลูกหลานหรือผู้ดูแล", "เจ้าหน้าที่หรือบุคลากรสุขภาพ"], "ท่านคือ")
    rich(doc, [("อายุ ........ ปี      วันที่ ................      รหัส ........", False, None)], size=14, after=4)
    h2(doc, "ส่วนที่ 1 · เรื่องที่เกิดขึ้นจริง")
    ticks(doc, ["เคย", "ไม่เคย", "จำไม่ได้"], "1. ในหนึ่งปีที่ผ่านมา ท่าน (หรือผู้สูงอายุที่ท่านดูแล) เคยล้มหรือเกือบล้มไหม")
    ticks(doc, ["เคย", "ไม่เคย"], "2. เคยมีใครวัดการทรงตัวหรือกำลังขาให้ไหม")
    ticks(doc, ["กังวลมาก", "กังวลบ้าง", "ไม่กังวล"], "3. ตอนนี้กังวลเรื่องการล้มแค่ไหน")
    ticks(doc, ["ตั้งเองได้", "ต้องมีคนช่วย", "ไม่ได้ใช้มือถือ"], "4. เวลาลงแอปใหม่ในมือถือ")
    h2(doc, "ส่วนที่ 2 · หลังอ่านคำอธิบายโปรแกรม")
    box(doc, "คำอธิบายโปรแกรม", SCRIPT_30S)
    rich(doc, [("5. จากที่อ่าน ท่านคิดว่าโปรแกรมนี้ทำอะไร (เขียนด้วยคำของท่านเอง)", True, None)], size=14, after=0)
    lines(doc, 2)
    ticks(doc, ["น่าจะช่วยได้", "ยังไม่แน่ใจ", "ไม่น่าช่วยได้"], "6. โปรแกรมนี้ช่วยเรื่องความเสี่ยงล้มได้ไหม")
    rich(doc, [("7. ถ้าจะใช้ที่บ้าน อะไรคืออุปสรรค", True, None)], size=14, after=0)
    lines(doc, 2)
    h2(doc, "ส่วนที่ 3 · เลือกสิ่งที่มีประโยชน์มากที่สุด 3 ข้อ ใส่เลข 1 2 3 หน้าข้อ")
    for fid, nm in FEATURES:
        rich(doc, [("[    ]  ", False, None), (fid + "  ", True, BRAND), (nm, False, None)], size=13.5, after=0, indent=0.3)
    para(doc, "", size=4)
    ticks(doc, ["ยินดีให้ทีมงานติดต่อกลับเพื่อมาลองใช้"])
    rich(doc, [("8. ข้อเสนอแนะอื่น", True, None)], size=14, after=0)
    lines(doc, 2)


def page_wrap(doc):
    h1(doc, "แบบสรุปรายคน", "กรอกภายใน 10 นาทีหลังจบการสัมภาษณ์ ตอนที่ยังจำได้ · หนึ่งแผ่นต่อหนึ่งคน")
    rich(doc, [("รหัส ........      กลุ่ม  ☐ ก  ☐ ข  ☐ ค      อายุ ........ ปี      วันที่ ................", False, None)], size=14, after=4)
    ticks(doc, ["เคยล้มหรือเกือบล้มใน 12 เดือน", "ไม่เคย", "ไม่ได้ถาม"])
    ticks(doc, ["ไม่เคยถูกประเมินการทรงตัวมาก่อน", "เคยถูกประเมิน", "ไม่ได้ถาม"])
    ticks(doc, ["อธิบายได้เองว่าโปรแกรมทำอะไร", "อธิบายไม่ได้"])
    ticks(doc, ["ตั้งแอปเองได้", "ต้องมีคนช่วยตั้ง", "ไม่ได้ใช้มือถือ"])
    ticks(doc, ["ให้ช่องทางติดต่อกลับไว้", "ไม่ให้"])
    rich(doc, [("บัตรฟีเจอร์ อันดับ 1 ........  2 ........  3 ........      ใบที่บอกว่าไม่จำเป็น ........", False, None)], size=14, after=4)
    for label in ["ข้อกังวลหลักของคนนี้ ประโยคเดียว", "คำพูดจริงที่ควรยกไปใส่ใบสรุป", "สิ่งที่เราจะเปลี่ยนในระบบเพราะคนนี้"]:
        rich(doc, [(label, True, None)], size=14, after=0)
        lines(doc, 2)
    h2(doc, "หลังกลับจากบ้านบางแค")
    for i, t in enumerate([
        "ถ่ายรูปแบบสัมภาษณ์และแบบสรุปรายคนที่กรอกแล้วทุกแผ่น (ไม่ต้องถ่ายใบยินยอมที่มีลายมือชื่อ)",
        "ส่งรูปให้ทีมถอดข้อมูลลงไฟล์ validation/results.json หนึ่งคนหนึ่งบล็อก",
        "สร้างใบสรุปหนึ่งหน้า ซึ่งคำนวณตัวเลขจากข้อมูลที่กรอกให้เอง แล้วอ่านทวนทั้งทีม",
        "ส่งแบบฟอร์มภารกิจก่อนเที่ยงวันที่ 16 กันยายน อย่ารอ 23.59 น.",
    ], 1):
        rich(doc, [(f"{i}.  ", True, BRAND), (t, False, None)], size=13.5, after=0, indent=0.3)


# ------------------------------------------------------------ ประกอบเล่ม
def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.3), Cm(1.3)
    sec.footer_distance = Cm(0.6)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn("w:cs"), FONT); st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    # เลขหน้าท้ายกระดาษ
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(fp.add_run("ชุดสัมภาษณ์ CareSignal · บ้านบางแค · หน้า "), 11, False, INK2)
    r = fp.add_run(); set_font(r, 11, False, INK2)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._element.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._element.append(it)
    for fn in (page_field_day, page_plan, page_consent, page_elder_a, page_elder_b,
               page_family, page_staff, page_cards, page_survey, page_wrap):
        fn(doc)
    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT_DOCX)
    return OUT_DOCX


def to_pdf(src, dst):
    """แปลงด้วย Microsoft Word ตัวจริง ให้ PDF ตรงกับที่ผู้ใช้เปิดใน Word ทุกบรรทัด"""
    import win32com.client as win32
    word = win32.DispatchEx("Word.Application"); word.Visible = False
    try:
        d = word.Documents.Open(os.path.abspath(src), ReadOnly=True)
        d.SaveAs2(os.path.abspath(dst), FileFormat=17)
        n = d.ComputeStatistics(2)
        d.Close(False)
        return n
    finally:
        word.Quit()


if __name__ == "__main__":
    out = build()
    print("เขียน", os.path.relpath(out, ROOT))
    if "--no-pdf" not in sys.argv:
        pages = to_pdf(OUT_DOCX, OUT_PDF)
        print("เขียน", os.path.relpath(OUT_PDF, ROOT), "·", pages, "หน้า")
